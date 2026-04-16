from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT_KR = "Malgun Gothic"
FONT_EN = "Aptos"


BLUE = RGBColor(33, 85, 214)
BLUE_DARK = RGBColor(37, 73, 153)
SLATE = RGBColor(85, 89, 112)
PURPLE = RGBColor(106, 70, 196)
PURPLE_DARK = RGBColor(85, 54, 167)
YELLOW = RGBColor(255, 232, 59)
RED = RGBColor(227, 61, 46)
ORANGE = RGBColor(230, 132, 30)
BLACK = RGBColor(17, 24, 39)
GRAY = RGBColor(99, 110, 128)
GRAY_LIGHT = RGBColor(230, 233, 240)
GRAY_LINE = RGBColor(201, 208, 220)
WHITE = RGBColor(255, 255, 255)


def box(slide, x, y, w, h, fill=None, line=GRAY_LINE, line_width=1.0, rounded=False):
    shape_type = (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)
    return shape


def add_textbox(
    slide,
    x,
    y,
    w,
    h,
    paragraphs,
    margin=0.08,
    valign=MSO_ANCHOR.TOP,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)

    for idx, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = para.get("align", PP_ALIGN.LEFT)
        p.space_before = Pt(0)
        p.space_after = Pt(para.get("space_after", 4))
        p.line_spacing = para.get("line_spacing", 1.15)

        if "segments" in para:
            segments = para["segments"]
        else:
            segments = [
                {
                    "text": para.get("text", ""),
                    "font_size": para.get("font_size", 18),
                    "bold": para.get("bold", False),
                    "color": para.get("color", BLACK),
                    "font_name": para.get("font_name", FONT_KR),
                }
            ]

        for seg in segments:
            run = p.add_run()
            run.text = seg["text"]
            font = run.font
            font.name = seg.get("font_name", FONT_KR)
            font.size = Pt(seg.get("font_size", 18))
            font.bold = seg.get("bold", False)
            font.italic = seg.get("italic", False)
            font.color.rgb = seg.get("color", BLACK)
    return shape


def add_title(slide, title):
    add_textbox(
        slide,
        0.22,
        0.08,
        4.0,
        0.44,
        [
            {
                "text": title,
                "font_size": 24,
                "bold": True,
                "color": BLUE,
                "space_after": 0,
            }
        ],
    )


def add_panel_frame(slide):
    box(slide, 0.12, 0.56, 4.42, 6.9, fill=None, line=GRAY_LINE, line_width=1.2)
    box(slide, 4.64, 0.56, 8.56, 6.9, fill=None, line=GRAY_LINE, line_width=1.2)


def add_panel_header(slide, x, y, w, text):
    box(slide, x, y, w, 0.44, fill=BLUE_DARK, line=BLUE_DARK)
    add_textbox(
        slide,
        x,
        y + 0.02,
        w,
        0.36,
        [
            {
                "text": text,
                "font_size": 17,
                "bold": True,
                "color": WHITE,
                "align": PP_ALIGN.CENTER,
                "space_after": 0,
            }
        ],
        margin=0.02,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_chip(slide, x, y, w, text):
    box(slide, x, y, w, 0.52, fill=PURPLE, line=PURPLE_DARK, rounded=True)
    add_textbox(
        slide,
        x,
        y + 0.03,
        w,
        0.44,
        [
            {
                "text": text,
                "font_size": 13.5,
                "bold": True,
                "color": WHITE,
                "align": PP_ALIGN.CENTER,
                "space_after": 0,
            }
        ],
        margin=0.02,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_dark_bar(slide, x, y, w, label):
    box(slide, x, y, w, 0.4, fill=SLATE, line=SLATE, rounded=True)
    add_textbox(
        slide,
        x + 0.1,
        y + 0.03,
        w - 0.2,
        0.3,
        [
            {
                "text": label,
                "font_size": 15.5,
                "bold": True,
                "color": YELLOW,
                "align": PP_ALIGN.LEFT,
                "space_after": 0,
            }
        ],
        margin=0,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_callout(slide, x, y, w, top_text, bottom_text):
    box(slide, x, y, w, 1.0, fill=BLACK, line=BLACK)
    add_textbox(
        slide,
        x + 0.14,
        y + 0.09,
        w - 0.28,
        0.34,
        [
            {
                "text": top_text,
                "font_size": 12.5,
                "bold": True,
                "color": WHITE,
                "align": PP_ALIGN.CENTER,
                "space_after": 2,
            },
            {
                "text": bottom_text,
                "font_size": 15,
                "bold": True,
                "color": YELLOW,
                "align": PP_ALIGN.CENTER,
                "space_after": 0,
            },
        ],
        margin=0,
    )


def add_footer(slide):
    add_textbox(
        slide,
        0.25,
        7.12,
        5.8,
        0.2,
        [
            {
                "text": "자료: deep-research-report.md 재구성",
                "font_size": 9,
                "color": GRAY,
                "space_after": 0,
            }
        ],
        margin=0,
    )


def add_bullet_lines(slide, x, y, w, items, font_size=14.5):
    paragraphs = []
    for text, color in items:
        paragraphs.append(
            {
                "segments": [
                    {
                        "text": "• ",
                        "font_size": font_size,
                        "bold": True,
                        "color": BLUE,
                    },
                    {
                        "text": text,
                        "font_size": font_size,
                        "color": color,
                    },
                ],
                "space_after": 5,
                "line_spacing": 1.2,
            }
        )
    add_textbox(slide, x, y, w, 2.0, paragraphs, margin=0)


def add_kpi_card(
    slide,
    x,
    y,
    w,
    title,
    desc,
    fill=GRAY_LIGHT,
    h=1.02,
    title_size=14,
    desc_size=11,
):
    box(slide, x, y, w, h, fill=fill, line=GRAY_LINE, line_width=1.0)
    add_textbox(
        slide,
        x + 0.08,
        y + 0.08,
        w - 0.16,
        h - 0.16,
        [
            {
                "text": title,
                "font_size": title_size,
                "bold": True,
                "color": BLUE_DARK,
                "align": PP_ALIGN.CENTER,
                "space_after": 2,
            },
            {
                "text": desc,
                "font_size": desc_size,
                "color": BLACK,
                "align": PP_ALIGN.CENTER,
                "space_after": 0,
            },
        ],
        margin=0.01,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_simple_bars(slide, x, y, w, h, labels, values, colors, title, max_value):
    box(slide, x, y, w, h, fill=WHITE, line=GRAY_LINE, line_width=1.0)
    add_textbox(
        slide,
        x,
        y + 0.02,
        w,
        0.25,
        [
            {
                "text": title,
                "font_size": 11.5,
                "bold": True,
                "color": BLACK,
                "align": PP_ALIGN.CENTER,
                "space_after": 0,
            }
        ],
        margin=0,
    )
    chart_left = x + 0.24
    chart_bottom = y + h - 0.24
    chart_top = y + 0.34
    usable_h = chart_bottom - chart_top
    bar_w = (w - 0.7) / max(len(values), 1)

    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(chart_left - 0.03),
        Inches(chart_bottom - 0.01),
        Inches(w - 0.42),
        Inches(0.01),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = GRAY
    line.line.fill.background()

    for idx, value in enumerate(values):
        bar_h = usable_h * (value / max_value)
        bx = chart_left + idx * bar_w + 0.11
        by = chart_bottom - bar_h
        box(slide, bx, by, 0.45, bar_h, fill=colors[idx], line=colors[idx])
        add_textbox(
            slide,
            bx - 0.08,
            chart_bottom + 0.02,
            0.62,
            0.22,
            [
                {
                    "text": labels[idx],
                    "font_size": 9.5,
                    "color": BLACK,
                    "align": PP_ALIGN.CENTER,
                    "space_after": 0,
                }
            ],
            margin=0,
        )
        add_textbox(
            slide,
            bx - 0.1,
            by - 0.18,
            0.7,
            0.18,
            [
                {
                    "text": f"{value:g}",
                    "font_size": 9,
                    "bold": True,
                    "color": colors[idx],
                    "align": PP_ALIGN.CENTER,
                    "space_after": 0,
                }
            ],
            margin=0,
        )


def add_ratio_chart(slide, x, y, w, h, first_label, first_value, second_label, second_value, title):
    box(slide, x, y, w, h, fill=WHITE, line=GRAY_LINE, line_width=1.0)
    add_textbox(
        slide,
        x,
        y + 0.02,
        w,
        0.25,
        [
            {
                "text": title,
                "font_size": 11.5,
                "bold": True,
                "color": BLACK,
                "align": PP_ALIGN.CENTER,
                "space_after": 0,
            }
        ],
        margin=0,
    )
    total_w = w - 0.48
    first_w = total_w * (first_value / (first_value + second_value))
    second_w = total_w - first_w
    bar_x = x + 0.24
    bar_y = y + 0.52
    bar_h = 0.42
    box(slide, bar_x, bar_y, first_w, bar_h, fill=BLUE, line=BLUE)
    box(slide, bar_x + first_w, bar_y, second_w, bar_h, fill=ORANGE, line=ORANGE)
    add_textbox(
        slide,
        bar_x,
        bar_y + 0.06,
        first_w,
        0.2,
        [
            {
                "text": f"{first_label} {first_value:.1f}%",
                "font_size": 10,
                "bold": True,
                "color": WHITE,
                "align": PP_ALIGN.CENTER,
                "space_after": 0,
            }
        ],
        margin=0,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_textbox(
        slide,
        bar_x + first_w,
        bar_y + 0.06,
        second_w,
        0.2,
        [
            {
                "text": f"{second_label} {second_value:.1f}%",
                "font_size": 10,
                "bold": True,
                "color": WHITE,
                "align": PP_ALIGN.CENTER,
                "space_after": 0,
            }
        ],
        margin=0,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_timeline_step(slide, x, y, label, title, desc, active=False):
    circle_color = BLUE if active else SLATE
    circle = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Inches(x + 0.4),
        Inches(y),
        Inches(0.42),
        Inches(0.42),
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = circle_color
    circle.line.color.rgb = circle_color
    add_textbox(
        slide,
        x + 0.4,
        y + 0.08,
        0.42,
        0.22,
        [
            {
                "text": label,
                "font_size": 10,
                "bold": True,
                "color": WHITE,
                "align": PP_ALIGN.CENTER,
                "space_after": 0,
            }
        ],
        margin=0,
    )
    if label != "5":
        connector = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            Inches(x + 0.8),
            Inches(y + 0.19),
            Inches(1.15),
            Inches(0.03),
        )
        connector.fill.solid()
        connector.fill.fore_color.rgb = GRAY_LINE
        connector.line.fill.background()

    fill = RGBColor(239, 244, 255) if active else WHITE
    line = BLUE if active else GRAY_LINE
    box(slide, x, y + 0.55, 1.95, 1.46, fill=fill, line=line, line_width=1.0)
    add_textbox(
        slide,
        x + 0.08,
        y + 0.63,
        1.79,
        1.28,
        [
            {
                "text": title,
                "font_size": 13.3,
                "bold": True,
                "color": BLUE_DARK if active else BLACK,
                "align": PP_ALIGN.CENTER,
                "space_after": 4,
            },
            {
                "text": desc,
                "font_size": 10.5,
                "color": BLACK,
                "align": PP_ALIGN.CENTER,
                "space_after": 0,
                "line_spacing": 1.15,
            },
        ],
        margin=0.03,
    )


def slide_policy_background(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "1) 정책적 배경")
    add_panel_frame(slide)
    add_panel_header(slide, 0.14, 0.58, 4.38, "정책 추진 필요성")
    add_panel_header(slide, 4.66, 0.58, 8.52, "법·제도 기반과 정책 과제")

    add_textbox(
        slide,
        0.42,
        1.02,
        3.82,
        1.5,
        [
            {
                "text": "석재산업은 이미",
                "font_size": 15.5,
                "color": BLACK,
                "align": PP_ALIGN.CENTER,
                "space_after": 2,
            },
            {
                "text": "계획·실태조사·진흥지구·인증·원산지 표시",
                "font_size": 17.5,
                "bold": True,
                "color": BLUE_DARK,
                "align": PP_ALIGN.CENTER,
                "space_after": 2,
            },
            {
                "text": "등 정책수단을 갖춘 분야",
                "font_size": 15.5,
                "color": BLACK,
                "align": PP_ALIGN.CENTER,
                "space_after": 8,
            },
            {
                "text": "이제 필요한 것은 흩어진 통계·행정·시장 데이터를 연결해 정책의 객관성을 높이는 일",
                "font_size": 14,
                "color": GRAY,
                "align": PP_ALIGN.CENTER,
                "line_spacing": 1.15,
                "space_after": 0,
            },
        ],
    )
    add_chip(slide, 0.38, 2.92, 1.05, "종합계획")
    add_chip(slide, 1.67, 2.92, 1.12, "실태조사")
    add_chip(slide, 3.03, 2.92, 1.1, "정책집행")
    add_callout(slide, 0.26, 5.98, 4.14, "<정책적 함의>", "기초 DB와 구조진단이 곧 정책 실행력")
    add_textbox(
        slide,
        0.38,
        6.99,
        4.0,
        0.3,
        [
            {
                "text": "“근거 기반 설계”로 지원·규제·인증 정책 정밀화",
                "font_size": 15,
                "bold": True,
                "color": RED,
                "align": PP_ALIGN.CENTER,
                "space_after": 0,
            }
        ],
        margin=0,
    )

    add_dark_bar(slide, 4.84, 1.12, 7.94, "(정책 여건) 법·제도는 이미 마련, 부족한 것은 데이터 정합성")
    add_bullet_lines(
        slide,
        5.06,
        1.7,
        7.5,
        [
            ("종합계획·시행계획 체계가 가동 중이며 정책은 집행 고도화 단계", BLACK),
            ("진흥지구 타당성조사는 경제효과·환경영향 등 객관 자료를 요구", BLACK),
            ("실태조사·원산지 표시·인증 운영을 뒷받침할 통합 DB가 필요", BLACK),
        ],
    )
    add_dark_bar(slide, 4.84, 3.36, 7.94, "(정책 과제) 공식통계·행정자료·시장거래 데이터를 하나로 결합")
    add_bullet_lines(
        slide,
        5.06,
        3.94,
        7.48,
        [
            ("시계열·분류 체계를 맞춰야 정책 효과를 정량 검증 가능", BLACK),
            ("공공조달, 품질인증, 원산지 표시 점검 결과를 함께 관리", BLACK),
            ("지원·규제·수출 전략을 동일한 근거체계 위에서 설계", BLACK),
        ],
    )
    add_kpi_card(slide, 5.02, 5.54, 2.25, "공식통계", "무역·생산·가격 동향")
    add_kpi_card(slide, 7.58, 5.54, 2.25, "행정자료", "허가·실태조사·점검 결과")
    add_kpi_card(slide, 10.14, 5.54, 2.25, "시장거래 데이터", "조달·유통·수요 구조")
    add_footer(slide)


def slide_industry_background(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "2) 산업적 배경")
    add_panel_frame(slide)
    add_panel_header(slide, 0.14, 0.58, 4.38, "시장 구조와 리스크")
    add_panel_header(slide, 4.66, 0.58, 8.52, "글로벌·국내 산업 현황")

    add_textbox(
        slide,
        0.32,
        1.0,
        4.02,
        1.82,
        [
            {
                "text": "석재산업",
                "font_size": 20,
                "bold": True,
                "color": BLUE_DARK,
                "align": PP_ALIGN.CENTER,
                "space_after": 3,
            },
            {
                "text": "세계 시장 변동성, 국내 수입 의존, 공공조달 중심 수요가 중첩된 산업",
                "font_size": 14.5,
                "color": BLACK,
                "align": PP_ALIGN.CENTER,
                "space_after": 3,
                "line_spacing": 1.2,
            },
            {
                "text": "생산·유통·수요를 함께 봐야 가격, 물류, 병목 원인을 설명 가능",
                "font_size": 14.2,
                "color": GRAY,
                "align": PP_ALIGN.CENTER,
                "line_spacing": 1.15,
                "space_after": 0,
            },
        ],
    )
    add_chip(slide, 0.38, 2.92, 1.05, "글로벌\n변동성")
    add_chip(slide, 1.67, 2.92, 1.12, "수입\n의존")
    add_chip(slide, 3.03, 2.92, 1.1, "공공조달\n수요")
    add_callout(slide, 0.26, 5.98, 4.14, "<산업적 함의>", "생산·유통·수요 전반의 입체 진단 필요")
    add_textbox(
        slide,
        0.34,
        6.99,
        4.08,
        0.3,
        [
            {
                "text": "수급·가격·물류 리스크를 함께 봐야 산업정책이 작동",
                "font_size": 14.2,
                "bold": True,
                "color": RED,
                "align": PP_ALIGN.CENTER,
                "space_after": 0,
            }
        ],
        margin=0,
    )

    add_dark_bar(slide, 4.84, 1.12, 7.94, "(글로벌) 2021년 생산 1.625억 톤, 수출 216.8억 달러, 가격 상승")
    add_bullet_lines(
        slide,
        5.06,
        1.72,
        7.5,
        [
            ("세계 생산량 1억 6,250만 톤으로 전년 대비 4.8% 증가", BLACK),
            ("세계 수출은 216.8억 달러 규모로 확대, 경쟁 구도 빠르게 변화", BLACK),
            ("석재 평균가격도 상승해 국내 수급과 가격 전략에 영향", BLACK),
        ],
    )
    add_dark_bar(slide, 4.84, 3.34, 7.94, "(국내) 수입 확대·폐기물 비중·공공조달 집중이 핵심 변수")
    add_bullet_lines(
        slide,
        5.06,
        3.94,
        7.5,
        [
            ("2021년 수입 7.589억 달러, 수출 197만 달러로 수입 의존 구조", BLACK),
            ("완제품 28.8% 대비 폐기물 71.2% 추정으로 자원효율 개선 여지 큼", BLACK),
            ("공공조달 천연석 거래는 2017~2021년 연평균 약 3,128억 원", BLACK),
        ],
    )
    add_simple_bars(
        slide,
        5.02,
        5.52,
        3.55,
        1.34,
        ["수입", "수출"],
        [758.9, 2.0],
        [BLUE, ORANGE],
        "<국내 무역 규모(백만 USD)>",
        800,
    )
    add_ratio_chart(
        slide,
        9.1,
        5.52,
        3.55,
        1.34,
        "완제품",
        28.8,
        "폐기물",
        71.2,
        "<채석량 대비 비중>",
    )
    add_footer(slide)


def slide_progress(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "3) 추진 경과")
    add_panel_frame(slide)
    add_panel_header(slide, 0.14, 0.58, 4.38, "제도 기반 구축 흐름")
    add_panel_header(slide, 4.66, 0.58, 8.52, "석재산업 추진 경과")

    add_textbox(
        slide,
        0.34,
        1.0,
        4.0,
        1.8,
        [
            {
                "text": "규제 중심 관리에서",
                "font_size": 15.5,
                "color": BLACK,
                "align": PP_ALIGN.CENTER,
                "space_after": 2,
            },
            {
                "text": "진흥·관리 체계로 전환",
                "font_size": 20,
                "bold": True,
                "color": BLUE_DARK,
                "align": PP_ALIGN.CENTER,
                "space_after": 4,
            },
            {
                "text": "현재는 종합계획과 현장 집행을 운영하면서 데이터 기반 고도화로 넘어가는 단계",
                "font_size": 14.2,
                "color": GRAY,
                "align": PP_ALIGN.CENTER,
                "line_spacing": 1.15,
                "space_after": 0,
            },
        ],
    )
    add_chip(slide, 0.38, 2.92, 1.05, "법 제정")
    add_chip(slide, 1.67, 2.92, 1.12, "계획 수립")
    add_chip(slide, 3.03, 2.92, 1.1, "현장 집행")
    add_callout(slide, 0.26, 5.98, 4.14, "<현재 단계>", "제도 구축 이후 정책 집행 고도화")
    add_textbox(
        slide,
        0.32,
        6.99,
        4.12,
        0.3,
        [
            {
                "text": "다음 단계는 통합 DB와 구조진단을 통한 정책 정밀화",
                "font_size": 14,
                "bold": True,
                "color": RED,
                "align": PP_ALIGN.CENTER,
                "space_after": 0,
            }
        ],
        margin=0,
    )

    add_dark_bar(slide, 4.84, 1.12, 7.94, "(추진 흐름) 법적 기반 확보 → 계획 가동 → 현장 집행 → 데이터 기반 고도화")
    add_timeline_step(slide, 4.92, 1.78, "1", "법적 기반 마련", "규제 중심에서\n진흥·관리 대상으로\n정책 범위 확장")
    add_timeline_step(slide, 6.95, 1.78, "2", "시행령·규칙 정비", "실태조사·인증·원산지\n표시 등 집행 근거 확보")
    add_timeline_step(slide, 8.98, 1.78, "3", "제1차 종합계획", "2022~2026 계획 공표,\n연차 시행계획 전제")
    add_timeline_step(slide, 11.01, 1.78, "4", "현장 집행 확대", "환경개선, 원산지 표시\n점검, 지자체 조례 연계")
    add_timeline_step(slide, 8.98, 3.92, "5", "고도화 과제", "기초 DB 구축과\n생산·유통·수요 구조진단", active=True)
    add_kpi_card(slide, 4.96, 6.02, 1.84, "재정", "근거 기반 지원", h=0.72, title_size=12.5, desc_size=9.2)
    add_kpi_card(slide, 6.95, 6.02, 1.84, "인프라", "거점·공동물류", h=0.72, title_size=12.5, desc_size=9.2)
    add_kpi_card(slide, 8.94, 6.02, 1.84, "인증", "품질·원산지 신뢰", h=0.72, title_size=12.5, desc_size=9.2)
    add_kpi_card(slide, 10.93, 6.02, 1.84, "수출지원", "HS·표준 대응", h=0.72, title_size=12.5, desc_size=9.2)
    add_footer(slide)


def main():
    output_dir = Path("output")
    output_dir.mkdir(exist_ok=True)
    output_path = output_dir / "석재산업_배경_3장_초안.pptx"

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide_policy_background(prs)
    slide_industry_background(prs)
    slide_progress(prs)
    prs.save(output_path)
    print(output_path.resolve())


if __name__ == "__main__":
    main()
