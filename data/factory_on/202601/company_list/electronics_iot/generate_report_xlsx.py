# -*- coding: utf-8 -*-
"""
표본할당표 대비 보유현황 집계표 엑셀 (전체)
- Sheet1: 전자산업 (할당/보유/여유)
- Sheet2: IoT가전 (할당/소계/확정/예상/확인필요)
  3색 구분: 할당(노랑) / 소계-확정-예상(초록) / 확인필요(주황)
"""
import sys, io
if sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

import pandas as pd
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

# ========================================
# 데이터 로드
# ========================================
BASE_DIR = r"d:\git_rk\data\factory_on\202601\company_list\electronics_iot"
INPUT_DIR = os.path.join(BASE_DIR, "260319_1204")
INPUT_FILE = os.path.join(INPUT_DIR, "company_electronics_iot_employees_10_and_over_F_260319_1204.csv")

print("[INFO] Loading CSV...")
df = pd.read_csv(INPUT_FILE, encoding='utf-8-sig')
print(f"[INFO] Total rows: {len(df)}")

size_order = ['10~29인', '30~99인', '100~299인', '300인+']

def map_size(s):
    s = str(s).strip()
    if s == '10~29인': return '10~29인'
    elif s == '30~99인': return '30~99인'
    elif s == '100~299인': return '100~299인'
    elif s in ['300~499인', '500~999인', '1000인 이상', '300인 이상', '300인+']: return '300인+'
    return None

df['종업원규모_그룹'] = df['종업원규모'].apply(map_size)

# ========================================
# 할당표
# ========================================
electronics_alloc = {
    '전자부품':         {'10~29인': 22, '30~99인': 13, '100~299인': 11, '300인+': 17},
    '컴퓨터주변기기':   {'10~29인': 12, '30~99인': 10, '100~299인': 4,  '300인+': 1},
    '방송통신장비':     {'10~29인': 16, '30~99인': 9,  '100~299인': 7,  '300인+': 5},
    '영상음향기기':     {'10~29인': 12, '30~99인': 10, '100~299인': 7,  '300인+': 1},
    '측정제어분석기기': {'10~29인': 28, '30~99인': 14, '100~299인': 10, '300인+': 8},
    '전기장비':         {'10~29인': 55, '30~99인': 23, '100~299인': 23, '300인+': 32},
}

iot_alloc = {
    '지능형 가전':             {'10~29인': 54, '30~99인': 25, '100~299인': 14, '300인+': 10},
    '홈헬스케어':              {'10~29인': 11, '30~99인': 7,  '100~299인': 3,  '300인+': 1},
    '홈네트워크 및 주거안전':  {'10~29인': 57, '30~99인': 20, '100~299인': 5,  '300인+': 3},
    '홈에너지':                {'10~29인': 28, '30~99인': 7,  '100~299인': 2,  '300인+': 3},
}

# ========================================
# 1. 전자산업 집계
# ========================================
print("[STEP 1] Electronics aggregation...")
df_elec = df[df['업종_전자산업'] == 'O'].copy()
elec_categories = ['전자부품', '컴퓨터주변기기', '방송통신장비', '영상음향기기', '측정제어분석기기', '전기장비']
elec_results = {}

for cat in elec_categories:
    subset = df_elec[df_elec['전자산업_대분류'] == cat]
    counts = {}
    for size in size_order:
        counts[size] = len(subset[subset['종업원규모_그룹'] == size])
    counts['합계'] = sum(counts.values())
    elec_results[cat] = counts

elec_total = {}
for size in size_order + ['합계']:
    elec_total[size] = sum(elec_results[cat][size] for cat in elec_categories)
elec_results['합계'] = elec_total

# ========================================
# 2. IoT가전 집계
# ========================================
print("[STEP 2] IoT aggregation...")
df_iot = df[df['가전IoT_예상'].isin(['확정(이전조사)', '확정(KEA추가)', '예상', '모호'])].copy()
print(f"[INFO] IoT related rows: {len(df_iot)}")

df_iot['is_confirmed'] = df_iot['가전IoT_예상'].str.contains('확정', na=False)
df_iot['is_expected'] = df_iot['가전IoT_예상'] == '예상'
df_iot['is_ambiguous'] = df_iot['가전IoT_예상'] == '모호'

def extract_iot_major(reason):
    if pd.isna(reason) or str(reason).strip() == '':
        return '미분류'
    reason = str(reason)
    for major in ['지능형 가전', '홈헬스케어', '홈네트워크 및 주거안전', '홈에너지']:
        if major in reason:
            return major
    return '미분류'

df_iot['iot_대분류'] = df_iot['예상_근거'].apply(extract_iot_major)

iot_categories = ['지능형 가전', '홈헬스케어', '홈네트워크 및 주거안전', '홈에너지']
iot_results = {}

for cat in iot_categories:
    subset = df_iot[df_iot['iot_대분류'] == cat]
    counts = {'확정': {}, '예상': {}, '모호': {}, '소계': {}}
    for size in size_order:
        ss = subset[subset['종업원규모_그룹'] == size]
        counts['확정'][size] = len(ss[ss['is_confirmed']])
        counts['예상'][size] = len(ss[ss['is_expected']])
        counts['모호'][size] = len(ss[ss['is_ambiguous']])
        counts['소계'][size] = counts['확정'][size] + counts['예상'][size]
    for key in counts:
        counts[key]['합계'] = sum(counts[key][s] for s in size_order)
    iot_results[cat] = counts

# 미분류
subset_unc = df_iot[df_iot['iot_대분류'] == '미분류']
iot_unc = {'확정': {}, '예상': {}, '모호': {}, '소계': {}}
for size in size_order:
    ss = subset_unc[subset_unc['종업원규모_그룹'] == size]
    iot_unc['확정'][size] = len(ss[ss['is_confirmed']])
    iot_unc['예상'][size] = len(ss[ss['is_expected']])
    iot_unc['모호'][size] = len(ss[ss['is_ambiguous']])
    iot_unc['소계'][size] = iot_unc['확정'][size] + iot_unc['예상'][size]
for key in iot_unc:
    iot_unc[key]['합계'] = sum(iot_unc[key][s] for s in size_order)

# 합계
iot_total = {'확정': {}, '예상': {}, '모호': {}, '소계': {}}
for size in size_order + ['합계']:
    for key in iot_total:
        val = sum(iot_results[cat][key].get(size, 0) for cat in iot_categories)
        val += iot_unc[key].get(size, 0)
        iot_total[key][size] = val
iot_results['합계'] = iot_total

# ========================================
# 엑셀 생성
# ========================================
print("[STEP 3] Generating Excel...")

wb = Workbook()

# --- 스타일 ---
thin = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
font_b = Font(name='맑은 고딕', size=10)
font_bold = Font(name='맑은 고딕', size=10, bold=True)
font_title = Font(name='맑은 고딕', size=12, bold=True)
font_h = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFF')
ca = Alignment(horizontal='center', vertical='center')
ra = Alignment(horizontal='right', vertical='center')
nf = '#,##0'
nf_diff = '+#,##0;-#,##0;0'

fill_hdr = PatternFill('solid', fgColor='4472C4')
fill_cat = PatternFill('solid', fgColor='D9E2F3')
fill_tot = PatternFill('solid', fgColor='B4C6E7')

# 색 1: 할당 (노랑)
fill_alloc = PatternFill('solid', fgColor='FFF2CC')
font_alloc = Font(name='맑은 고딕', size=10, bold=True, color='7F6000')

# 색 2: 소계/확정/예상 (초록)
fill_subtotal = PatternFill('solid', fgColor='C6EFCE')
font_subtotal = Font(name='맑은 고딕', size=10, bold=True, color='006100')
fill_detail = PatternFill('solid', fgColor='E2EFDA')

# 색 3: 확인필요 (주황)
fill_check = PatternFill('solid', fgColor='FCE4D6')
font_check = Font(name='맑은 고딕', size=10, color='BF4000')
font_check_b = Font(name='맑은 고딕', size=10, bold=True, color='BF4000')

# 전자산업용: 보유/여유
fill_hold = PatternFill('solid', fgColor='E2EFDA')
fill_marg_p = PatternFill('solid', fgColor='DDEBF7')
fill_marg_n = PatternFill('solid', fgColor='FCE4EC')

def cell(ws, r, c, val, font=font_b, fill=None, align=ra, fmt=None):
    cl = ws.cell(row=r, column=c, value=val)
    cl.font = font; cl.alignment = align; cl.border = thin
    if fill: cl.fill = fill
    if fmt: cl.number_format = fmt
    return cl

def write_header(ws, row):
    for i, v in enumerate(['', '구분', '10~29인', '30~99인', '100~299인', '300인+', '합계']):
        cell(ws, row, i+1, v, font=font_h, fill=fill_hdr, align=ca)

# ========================================
# Sheet 1: 전자산업
# ========================================
ws1 = wb.active
ws1.title = '전자산업'

ws1.merge_cells('A1:G1')
cell(ws1, 1, 1, '전자산업 할당표 대비 보유현황', font=font_title, align=ca)

write_header(ws1, 3)
ws1.cell(row=3, column=1).value = '대분류'

row = 4
nums = {'전자부품': '(1)', '컴퓨터주변기기': '(2)', '방송통신장비': '(3)',
        '영상음향기기': '(4)', '측정제어분석기기': '(5)', '전기장비': '(6)'}

for cat in elec_categories + ['합계']:
    is_tot = (cat == '합계')
    hold = elec_results[cat]

    if is_tot:
        label = '합계'
        av = {s: sum(electronics_alloc[c_][s] for c_ in elec_categories) for s in size_order}
        av['합계'] = sum(av.values())
        cf, bg = font_bold, fill_tot
    else:
        label = f"{nums[cat]} {cat}"
        av = dict(electronics_alloc[cat])
        av['합계'] = sum(av.values())
        cf, bg = font_b, fill_cat

    ws1.merge_cells(start_row=row, start_column=1, end_row=row+2, end_column=1)
    cell(ws1, row, 1, label, font=cf, fill=bg, align=ca)
    for rr in range(row, row+3):
        ws1.cell(row=rr, column=1).border = thin

    # 할당
    cell(ws1, row, 2, '할당', font=font_alloc, fill=fill_alloc, align=ca)
    for i, s in enumerate(size_order):
        cell(ws1, row, 3+i, av[s], font=font_alloc, fill=fill_alloc, fmt=nf)
    cell(ws1, row, 7, av['합계'], font=font_alloc, fill=fill_alloc, fmt=nf)

    # 보유
    cell(ws1, row+1, 2, '보유', font=cf, fill=fill_hold, align=ca)
    for i, s in enumerate(size_order):
        cell(ws1, row+1, 3+i, hold[s], font=cf, fill=fill_hold, fmt=nf)
    cell(ws1, row+1, 7, hold['합계'], font=font_bold, fill=fill_hold, fmt=nf)

    # 여유
    cell(ws1, row+2, 2, '여유', font=cf, fill=fill_marg_p, align=ca)
    for i, s in enumerate(size_order):
        d = hold[s] - av[s]
        f = fill_marg_p if d >= 0 else fill_marg_n
        cell(ws1, row+2, 3+i, d, font=cf, fill=f, fmt=nf_diff)
    dt = hold['합계'] - av['합계']
    ft = fill_marg_p if dt >= 0 else fill_marg_n
    cell(ws1, row+2, 7, dt, font=font_bold, fill=ft, fmt=nf_diff)

    row += 3

ws1.column_dimensions['A'].width = 22
ws1.column_dimensions['B'].width = 8
for cl in ['C', 'D', 'E', 'F', 'G']:
    ws1.column_dimensions[cl].width = 12

# ========================================
# Sheet 2: IoT가전
# ========================================
ws2 = wb.create_sheet('IoT가전')

ws2.merge_cells('A1:G1')
cell(ws2, 1, 1, 'IoT가전 할당표 대비 보유현황', font=font_title, align=ca)

write_header(ws2, 3)
ws2.cell(row=3, column=1).value = '세부 업종'

row = 4
for cat in iot_categories + ['합계']:
    is_tot = (cat == '합계')

    if is_tot:
        label = '합계'
        av = {s: sum(iot_alloc[c_][s] for c_ in iot_categories) for s in size_order}
        av['합계'] = sum(av.values())
        r = iot_results['합계']
        cf, bg = font_bold, fill_tot
    else:
        label = cat
        av = dict(iot_alloc[cat])
        av['합계'] = sum(iot_alloc[cat].values())
        r = iot_results[cat]
        cf, bg = font_b, fill_cat

    # 카테고리 (5행 병합)
    ws2.merge_cells(start_row=row, start_column=1, end_row=row+4, end_column=1)
    cell(ws2, row, 1, label, font=cf, fill=bg, align=ca)
    for rr in range(row, row+5):
        ws2.cell(row=rr, column=1).border = thin

    # 행 0: 할당 (색 1 - 노랑)
    cell(ws2, row, 2, '할당', font=font_alloc, fill=fill_alloc, align=ca)
    for i, s in enumerate(size_order):
        cell(ws2, row, 3+i, av[s], font=font_alloc, fill=fill_alloc, fmt=nf)
    cell(ws2, row, 7, av['합계'], font=font_alloc, fill=fill_alloc, fmt=nf)

    # 행 1: 소계 (색 2 진하게 - 초록)
    cell(ws2, row+1, 2, '소계', font=font_subtotal, fill=fill_subtotal, align=ca)
    for i, s in enumerate(size_order):
        cell(ws2, row+1, 3+i, r['소계'][s], font=font_subtotal, fill=fill_subtotal, fmt=nf)
    cell(ws2, row+1, 7, r['소계']['합계'], font=font_subtotal, fill=fill_subtotal, fmt=nf)

    # 행 2: 확정 (색 2 연하게)
    cell(ws2, row+2, 2, '  확정', font=font_b, fill=fill_detail, align=ca)
    for i, s in enumerate(size_order):
        cell(ws2, row+2, 3+i, r['확정'][s], font=font_b, fill=fill_detail, fmt=nf)
    cell(ws2, row+2, 7, r['확정']['합계'], font=font_bold, fill=fill_detail, fmt=nf)

    # 행 3: 예상 (색 2 연하게)
    cell(ws2, row+3, 2, '  예상', font=font_b, fill=fill_detail, align=ca)
    for i, s in enumerate(size_order):
        cell(ws2, row+3, 3+i, r['예상'][s], font=font_b, fill=fill_detail, fmt=nf)
    cell(ws2, row+3, 7, r['예상']['합계'], font=font_bold, fill=fill_detail, fmt=nf)

    # 행 4: 확인필요 (색 3 - 주황)
    cell(ws2, row+4, 2, '확인필요', font=font_check_b, fill=fill_check, align=ca)
    for i, s in enumerate(size_order):
        cell(ws2, row+4, 3+i, r['모호'][s], font=font_check, fill=fill_check, fmt=nf)
    cell(ws2, row+4, 7, r['모호']['합계'], font=font_check_b, fill=fill_check, fmt=nf)

    row += 5

# 주석
row += 1
ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
note1 = f"* 미분류(대분류 미확인): 확정 {iot_unc['확정']['합계']}건, 예상 {iot_unc['예상']['합계']}건, 확인필요 {iot_unc['모호']['합계']}건 - 위 합계에 포함"
ws2.cell(row=row, column=1, value=note1).font = Font(name='맑은 고딕', size=9, italic=True, color='666666')

row += 1
ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
note2 = "* 소계 = 확정 + 예상 | 확인필요 = 추가 검토 필요 업체"
ws2.cell(row=row, column=1, value=note2).font = Font(name='맑은 고딕', size=9, italic=True, color='666666')

ws2.column_dimensions['A'].width = 26
ws2.column_dimensions['B'].width = 10
for cl in ['C', 'D', 'E', 'F', 'G']:
    ws2.column_dimensions[cl].width = 12

# ========================================
# 저장
# ========================================
output_path = os.path.join(INPUT_DIR, "allocation_vs_holdings_report.xlsx")
wb.save(output_path)
print(f"\n[INFO] Excel saved to: {output_path}")
print("[DONE]")
