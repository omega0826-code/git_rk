# -*- coding: utf-8 -*-
"""
표본할당표 대비 보유현황 집계표 엑셀 - 샘플
IoT가전: 할당 / 소계 / 확정 / 예상 / 확인필요(모호)
3색 구분 적용
"""
import sys, io
if sys.stdout.encoding != 'utf-8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

import pandas as pd
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

# ========================================
# 데이터 로드
# ========================================
BASE_DIR = r"d:\git_rk\data\factory_on\202601\company_list\electronics_iot"
INPUT_DIR = os.path.join(BASE_DIR, "260319_1204")
INPUT_FILE = os.path.join(INPUT_DIR, "company_electronics_iot_employees_10_and_over_F_260319_1204.csv")

print("[INFO] Loading CSV...")
df = pd.read_csv(INPUT_FILE, encoding='utf-8-sig')
print(f"[INFO] Total rows: {len(df)}")

# ========================================
# 종업원규모 매핑
# ========================================
size_order = ['10~29인', '30~99인', '100~299인', '300인+']

def map_size(s):
    s = str(s).strip()
    if s == '10~29인': return '10~29인'
    elif s == '30~99인': return '30~99인'
    elif s == '100~299인': return '100~299인'
    elif s in ['300~499인', '500~999인', '1000인 이상', '300인 이상', '300인+']: return '300인+'
    return None

df['종업원규모_그룹'] = df['종업원규모'].apply(map_size)

# ========================================
# 할당표
# ========================================
iot_alloc = {
    '지능형 가전':             {'10~29인': 54, '30~99인': 25, '100~299인': 14, '300인+': 10},
    '홈헬스케어':              {'10~29인': 11, '30~99인': 7,  '100~299인': 3,  '300인+': 1},
    '홈네트워크 및 주거안전':  {'10~29인': 57, '30~99인': 20, '100~299인': 5,  '300인+': 3},
    '홈에너지':                {'10~29인': 28, '30~99인': 7,  '100~299인': 2,  '300인+': 3},
}

# ========================================
# IoT가전 집계 (확정/예상/모호 모두 포함)
# ========================================
print("[STEP] IoT aggregation...")

df_iot = df[df['가전IoT_예상'].isin(['확정(이전조사)', '확정(KEA추가)', '예상', '모호'])].copy()
print(f"[INFO] IoT related rows: {len(df_iot)}")

df_iot['is_confirmed'] = df_iot['가전IoT_예상'].str.contains('확정', na=False)
df_iot['is_expected'] = df_iot['가전IoT_예상'] == '예상'
df_iot['is_ambiguous'] = df_iot['가전IoT_예상'] == '모호'

# IoT 대분류 추출
def extract_iot_major(reason):
    if pd.isna(reason) or str(reason).strip() == '':
        return '미분류'
    reason = str(reason)
    for major in ['지능형 가전', '홈헬스케어', '홈네트워크 및 주거안전', '홈에너지']:
        if major in reason:
            return major
    return '미분류'

df_iot['iot_대분류'] = df_iot['예상_근거'].apply(extract_iot_major)

iot_categories = ['지능형 가전', '홈헬스케어', '홈네트워크 및 주거안전', '홈에너지']
iot_results = {}

for cat in iot_categories:
    subset = df_iot[df_iot['iot_대분류'] == cat]
    counts = {'확정': {}, '예상': {}, '모호': {}, '소계': {}}
    for size in size_order:
        ss = subset[subset['종업원규모_그룹'] == size]
        counts['확정'][size] = len(ss[ss['is_confirmed']])
        counts['예상'][size] = len(ss[ss['is_expected']])
        counts['모호'][size] = len(ss[ss['is_ambiguous']])
        counts['소계'][size] = counts['확정'][size] + counts['예상'][size]
    for key in counts:
        counts[key]['합계'] = sum(counts[key][s] for s in size_order)
    iot_results[cat] = counts

# 미분류
subset_unc = df_iot[df_iot['iot_대분류'] == '미분류']
iot_unc = {'확정': {}, '예상': {}, '모호': {}, '소계': {}}
for size in size_order:
    ss = subset_unc[subset_unc['종업원규모_그룹'] == size]
    iot_unc['확정'][size] = len(ss[ss['is_confirmed']])
    iot_unc['예상'][size] = len(ss[ss['is_expected']])
    iot_unc['모호'][size] = len(ss[ss['is_ambiguous']])
    iot_unc['소계'][size] = iot_unc['확정'][size] + iot_unc['예상'][size]
for key in iot_unc:
    iot_unc[key]['합계'] = sum(iot_unc[key][s] for s in size_order)

# 합계
iot_total = {'확정': {}, '예상': {}, '모호': {}, '소계': {}}
for size in size_order + ['합계']:
    for key in iot_total:
        val = sum(iot_results[cat][key].get(size, 0) for cat in iot_categories)
        val += iot_unc[key].get(size, 0)
        iot_total[key][size] = val
iot_results['합계'] = iot_total

# ========================================
# 엑셀 생성 (샘플: 지능형 가전 + 합계만)
# ========================================
print("[STEP] Generating sample Excel...")

wb = Workbook()

# --- 스타일 정의 ---
thin = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
font_h = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFF')
font_b = Font(name='맑은 고딕', size=10)
font_bold = Font(name='맑은 고딕', size=10, bold=True)
font_title = Font(name='맑은 고딕', size=12, bold=True)
ca = Alignment(horizontal='center', vertical='center')
ra = Alignment(horizontal='right', vertical='center')
nf = '#,##0'

# 헤더 (진한 파랑)
fill_hdr = PatternFill('solid', fgColor='4472C4')
# 카테고리 배경
fill_cat = PatternFill('solid', fgColor='D9E2F3')
# 합계 행 배경
fill_tot = PatternFill('solid', fgColor='B4C6E7')

# 색 1: 할당 (노란 계열)
fill_alloc = PatternFill('solid', fgColor='FFF2CC')
font_alloc = Font(name='맑은 고딕', size=10, bold=True, color='7F6000')

# 색 2: 소계/확정/예상 (초록 계열)
fill_subtotal = PatternFill('solid', fgColor='C6EFCE')
font_subtotal = Font(name='맑은 고딕', size=10, bold=True, color='006100')
fill_confirmed = PatternFill('solid', fgColor='E2EFDA')
fill_expected = PatternFill('solid', fgColor='E2EFDA')

# 색 3: 확인필요 (주황/살구 계열 - 구분 명확)
fill_check = PatternFill('solid', fgColor='FCE4D6')
font_check = Font(name='맑은 고딕', size=10, color='BF4000')
font_check_bold = Font(name='맑은 고딕', size=10, bold=True, color='BF4000')

def cell(ws, r, c, val, font=font_b, fill=None, align=ra, fmt=None):
    cl = ws.cell(row=r, column=c, value=val)
    cl.font = font
    cl.alignment = align
    cl.border = thin
    if fill: cl.fill = fill
    if fmt: cl.number_format = fmt
    return cl

# ========================================
# IoT가전 시트
# ========================================
ws = wb.active
ws.title = 'IoT가전'

# 타이틀
ws.merge_cells('A1:G1')
cell(ws, 1, 1, 'IoT가전 할당표 대비 보유현황', font=font_title, align=ca)

# 헤더
for i, v in enumerate(['세부 업종', '구분', '10~29인', '30~99인', '100~299인', '300인+', '합계']):
    cell(ws, 3, i+1, v, font=font_h, fill=fill_hdr, align=ca)

# 데이터 행 (샘플: 지능형가전 + 합계)
sample_cats = ['지능형 가전', '합계']
row = 4

for cat in sample_cats:
    is_tot = (cat == '합계')
    
    if is_tot:
        alloc_v = {s: sum(iot_alloc[c_][s] for c_ in iot_categories) for s in size_order}
        alloc_v['합계'] = sum(alloc_v.values())
        r = iot_results['합계']
        c_font = font_bold
        c_fill = fill_tot
    else:
        alloc_v = dict(iot_alloc[cat])
        alloc_v['합계'] = sum(iot_alloc[cat].values())
        r = iot_results[cat]
        c_font = font_b
        c_fill = fill_cat

    # 카테고리 셀 (5행 병합)
    ws.merge_cells(start_row=row, start_column=1, end_row=row+4, end_column=1)
    cell(ws, row, 1, cat, font=c_font, fill=c_fill, align=ca)
    for rr in range(row, row+5):
        ws.cell(row=rr, column=1).border = thin

    # 행 0: 할당 (색 1)
    cell(ws, row, 2, '할당', font=font_alloc, fill=fill_alloc, align=ca)
    for i, s in enumerate(size_order):
        cell(ws, row, 3+i, alloc_v[s], font=font_alloc, fill=fill_alloc, fmt=nf)
    cell(ws, row, 7, alloc_v['합계'], font=font_alloc, fill=fill_alloc, fmt=nf)

    # 행 1: 소계 (색 2 진하게)
    cell(ws, row+1, 2, '소계', font=font_subtotal, fill=fill_subtotal, align=ca)
    for i, s in enumerate(size_order):
        cell(ws, row+1, 3+i, r['소계'][s], font=font_subtotal, fill=fill_subtotal, fmt=nf)
    cell(ws, row+1, 7, r['소계']['합계'], font=font_subtotal, fill=fill_subtotal, fmt=nf)

    # 행 2: 확정 (색 2 연하게)
    cell(ws, row+2, 2, '  확정', font=font_b, fill=fill_confirmed, align=ca)
    for i, s in enumerate(size_order):
        cell(ws, row+2, 3+i, r['확정'][s], font=font_b, fill=fill_confirmed, fmt=nf)
    cell(ws, row+2, 7, r['확정']['합계'], font=font_bold, fill=fill_confirmed, fmt=nf)

    # 행 3: 예상 (색 2 연하게)
    cell(ws, row+3, 2, '  예상', font=font_b, fill=fill_expected, align=ca)
    for i, s in enumerate(size_order):
        cell(ws, row+3, 3+i, r['예상'][s], font=font_b, fill=fill_expected, fmt=nf)
    cell(ws, row+3, 7, r['예상']['합계'], font=font_bold, fill=fill_expected, fmt=nf)

    # 행 4: 확인필요 (색 3)
    cell(ws, row+4, 2, '확인필요', font=font_check_bold, fill=fill_check, align=ca)
    for i, s in enumerate(size_order):
        cell(ws, row+4, 3+i, r['모호'][s], font=font_check, fill=fill_check, fmt=nf)
    cell(ws, row+4, 7, r['모호']['합계'], font=font_check_bold, fill=fill_check, fmt=nf)

    row += 5

# 주석
row += 1
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
ws.cell(row=row, column=1, value="* 소계 = 확정 + 예상 | 확인필요 = 기존 '모호' 분류 (추가 검토 필요)").font = Font(name='맑은 고딕', size=9, italic=True, color='666666')

# 열 너비
ws.column_dimensions['A'].width = 26
ws.column_dimensions['B'].width = 10
for cl in ['C', 'D', 'E', 'F', 'G']:
    ws.column_dimensions[cl].width = 12

# 저장
sample_path = os.path.join(INPUT_DIR, "allocation_vs_holdings_sample.xlsx")
wb.save(sample_path)
print(f"\n[INFO] Sample saved to: {sample_path}")
print("[DONE]")
